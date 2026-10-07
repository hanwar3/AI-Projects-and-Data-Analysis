"""Offline tests for the guarantees the spec calls non-negotiable."""
import sqlite3
from pathlib import Path

import pytest

import config
import db
import monitor
from matching import norm_date, score_notice
from sources import PartialFetch

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def temp_db(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(config, "LOG_PATH", tmp_path / "test.log")
    monkeypatch.setattr(config, "NOTIFY_MIN_SCORE", None)


def row(ref, title="Consultancy for a flood early warning system", deadline="2099-01-01"):
    return {"source": "Test", "type": "RFP", "title": title, "detail": "", "country": "Nowhere",
            "deadline": deadline, "ref": ref, "url": f"https://example.org/{ref}"}


def fake(name, rows=(), exc=None):
    def fetch():
        if exc:
            raise exc
        return list(rows)
    fetch.__name__, fetch.label = f"fetch_{name}", name
    return fetch


def test_repeat_run_is_a_no_op_and_keeps_date_found():
    src = fake("a", [row("1"), row("2"), row("2")])
    first = monitor.run_cycle(trigger="scheduled", sources=[src])
    stamps = {r["notice_id"]: r["date_found"] for r in db.all_notices()}
    second = monitor.run_cycle(trigger="manual", sources=[src])
    assert (first["new"], second["new"]) == (2, 0)
    assert {r["notice_id"]: r["date_found"] for r in db.all_notices()} == stamps


def test_duplicates_are_blocked_by_the_database_itself():
    db.init_db()
    notice_id = monitor.prepare([row("x")])[0]["notice_id"]
    db.insert_notices(monitor.prepare([row("x")]))
    con = sqlite3.connect(config.DB_PATH)
    with pytest.raises(sqlite3.IntegrityError):
        con.execute("INSERT INTO notices (notice_id, date_found) VALUES (?, '2026-01-01')", (notice_id,))
    con.close()


def test_every_source_is_logged_including_failures_and_zero_new():
    monitor.run_cycle(sources=[fake("ok", [row("1")]), fake("quiet"),
                               fake("boom", exc=ConnectionError("site down"))])
    runs = {r["source"]: r for r in db.recent_runs()}
    assert (runs["ok"]["status"], runs["ok"]["new_count"]) == ("ok", 1)
    assert (runs["quiet"]["status"], runs["quiet"]["new_count"]) == ("ok", 0)
    assert runs["boom"]["status"] == "error" and "site down" in runs["boom"]["error"]


def test_partial_fetch_keeps_rows_and_records_the_gap():
    monitor.run_cycle(sources=[fake("half", exc=PartialFetch([row("1")], "1/2 queries failed"))])
    run = db.recent_runs()[0]
    assert (run["status"], run["new_count"]) == ("partial", 1) and "1/2" in run["error"]


def test_keyword_gate_expiry_and_accent_folding():
    kept = monitor.prepare([row("a", title="Supply of office furniture"),
                            row("b", deadline="2001-01-01"),
                            row("c"),
                            row("d", title="Évaluation de la résilience urbaine")], today="2026-09-11")
    assert [r["ref"] for r in kept] == ["c", "d"]


def test_untrusted_urls_are_dropped():
    bad = row("e")
    bad["url"] = "javascript:alert(1)"
    assert monitor.prepare([bad])[0]["url"] == ""


@pytest.mark.parametrize("raw, iso", [
    ("09-Oct-26 06:00 AM (New York time)", "2026-10-09"),
    ("21-Sep-2026", "2026-09-21"),
    ("9/20/26 11:59 PM", "2026-09-20"),
    ("Friday, January 09, 2026", "2026-01-09"),
    ("2026-09-18T03:59:00.000+0000", "2026-09-18"),
    ("Thu, 10 Sep 2026 16:36:10 +0000", "2026-09-10"),
    ("14 September 2026", "2026-09-14"),
    ("no deadline: rolling application process", ""),
])
def test_norm_date(raw, iso):
    assert norm_date(raw) == iso


def test_score_ranks_field_consultancy_over_side_matches_and_goods():
    strong, _ = score_notice("Consultancy for a disaster risk reduction strategy")
    side, _ = score_notice("Customs clearing services", "Integrated Flood Resilience and Adaptation Project")
    goods, _ = score_notice("Supply of vehicles", "Flood Resilience Project")
    assert strong > side > goods


def test_exports_open_cleanly_and_neutralise_formulas(tmp_path):
    import pandas as pd
    from openpyxl import load_workbook

    import exports
    frame = exports.spreadsheet_safe(pd.DataFrame(
        [{"title": '=HYPERLINK("http://evil")', "detail": "Résilience – côtière", "score": 57}]))
    path = tmp_path / "view.xlsx"
    path.write_bytes(exports.to_xlsx(frame))
    ws = load_workbook(path).active
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == ("'=HYPERLINK(\"http://evil\")", "Résilience – côtière", 57)
    assert exports.to_csv(frame).startswith(b"\xef\xbb\xbf")


def test_scheduler_button_and_cli_share_one_ingestion_path():
    writers = sorted(p.name for p in ROOT.glob("*.py")
                     if p.name != "db.py" and "insert_notices(" in p.read_text(encoding="utf-8"))
    assert writers == ["monitor.py"]
    app = (ROOT / "app.py").read_text(encoding="utf-8")
    assert 'run_cycle(trigger="manual"' in app          # the Refresh button
    assert "apscheduler" not in app.lower()               # nothing runs on a timer
